"""Build the editable presentation and Word/PDF report from project artifacts."""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
RESULTS = ROOT / "results"
METRICS = pd.read_csv(RESULTS / "metrics.csv")

INK = "213735"
GREEN = "2D6258"
DARK = "1C3935"
CREAM = "F5F4ED"
WHITE = "FFFEFA"
MUTED = "70837D"
SAGE = "DCE9DF"
LINE = "D9E1DA"
AMBER = "EAB45E"
CORAL = "D97754"
PALE = "EDF2EB"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_rect(slide, x, y, w, h, fill, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background() if line is None else None
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    return s


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.01); tf.margin_right = Inches(0.01)
    tf.margin_top = Inches(0.01); tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return box


def add_bullets(slide, items, x, y, w, h, size=16, color=INK, gap=10):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(.02); tf.margin_top = Inches(.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
    return box


def add_footer(slide, n, dark=False):
    color = "B8CCC3" if dark else MUTED
    add_text(slide, "MOSAIC-AIR  ·  ENVIRONMENTAL MANAGEMENT", .55, 7.14, 7, .18, 8, color, True)
    add_text(slide, f"{n:02d}", 12.2, 7.1, .55, .22, 9, color, True, align=PP_ALIGN.RIGHT)


def new_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(DARK if dark else CREAM)
    return slide


def section_title(slide, eyebrow, title, subtitle=None, dark=False):
    base = "E5B96C" if dark else GREEN
    color = WHITE if dark else INK
    add_text(slide, eyebrow.upper(), .65, .42, 11.9, .26, 10, base, True)
    add_text(slide, title, .65, .79, 12, .65, 27, color, True, font="Aptos Display")
    if subtitle:
        add_text(slide, subtitle, .68, 1.48, 11.8, .38, 13, "C4D2CC" if dark else MUTED)


def add_pill(slide, text, x, y, w, fill=SAGE, color=GREEN):
    add_rect(slide, x, y, w, .32, fill, True)
    add_text(slide, text, x + .08, y + .055, w - .16, .21, 9, color, True, align=PP_ALIGN.CENTER)


def make_slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "MOSAIC-Air: Wind-Aware Urban Air Pollution Forecasting"
    prs.core_properties.subject = "Environmental Management project presentation"
    prs.core_properties.author = "Aakash, Shaik Mohammed Atif, Shivadev"

    # 1 · Cover
    s = new_slide(prs, dark=True)
    # atmospheric network graphic
    for x, y, size in [(9.2, 1.5, .13), (10.7, 1.0, .19), (11.7, 2.1, .15), (10.0, 3.0, .24), (12.1, 4.0, .16), (8.9, 4.5, .17), (11.1, 5.4, .12), (9.8, 6.0, .12)]:
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        circ.fill.solid(); circ.fill.fore_color.rgb = rgb(AMBER); circ.line.fill.background()
    for x1, y1, x2, y2 in [(9.3,1.6,10.8,1.1),(10.8,1.1,11.8,2.2),(11.8,2.2,10.1,3.1),(10.1,3.1,12.2,4.1),(10.1,3.1,8.95,4.6),(8.95,4.6,9.85,6.1),(9.85,6.1,11.15,5.5),(11.15,5.5,12.2,4.1),(10.8,1.1,10.1,3.1)]:
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = rgb("52746A"); line.line.width = Pt(1.3)
    add_pill(s, "ENVIRONMENTAL MANAGEMENT · 22ENV231", .7, .55, 3.8, fill="31564D", color="D5E6DB")
    add_text(s, "MOSAIC-Air", .7, 1.62, 7.5, .95, 42, WHITE, True, font="Aptos Display")
    add_text(s, "Wind-aware urban air\npollution forecasting", .73, 2.68, 7.7, 1.25, 26, "D7E5DD", False, font="Aptos Display")
    add_text(s, "A spatiotemporal deep-learning pilot for short-term PM₂.₅ forecasts, emerging hotspots, and forecast uncertainty.", .76, 4.24, 6.8, .78, 15, "B8CCC3")
    add_text(s, "Aakash  ·  Shaik Mohammed Atif  ·  Shivadev", .77, 6.18, 7.2, .3, 13, WHITE, True)
    add_text(s, "Amrita Vishwa Vidyapeetham · Amritapuri Campus", .77, 6.53, 7.2, .23, 10, "B8CCC3")
    add_footer(s, 1, True)

    # 2 · Problem
    s = new_slide(prs)
    section_title(s, "01 / Motivation", "Air risk changes across both time and place", "A single citywide number can hide where the next local episode is forming.")
    cards = [
        ("01", "Uneven monitoring", "Sparse station coverage leaves gaps between measured locations."),
        ("02", "Moving conditions", "Wind and weather change how pollution builds and moves through a city."),
        ("03", "Action needs context", "Managers need forecasts with uncertainty and a clear location, not a score alone."),
    ]
    for i, (num, title, body) in enumerate(cards):
        x = .7 + i * 4.18
        add_rect(s, x, 2.25, 3.82, 2.4, WHITE, True, LINE)
        add_text(s, num, x+.22, 2.49, .5, .3, 12, CORAL, True)
        add_text(s, title, x+.22, 3.02, 3.32, .38, 19, INK, True)
        add_text(s, body, x+.22, 3.54, 3.3, .83, 14, MUTED)
    add_rect(s, .7, 5.1, 11.98, 1.12, SAGE, True)
    add_text(s, "PROJECT QUESTION", .98, 5.37, 1.9, .23, 10, GREEN, True)
    add_text(s, "Can wind-conditioned station relationships improve multi-horizon PM₂.₅ forecasts and help flag emerging high-pollution episodes?", 2.95, 5.27, 9.2, .57, 17, INK, True)
    add_footer(s, 2)

    # 3 · Objectives
    s = new_slide(prs, dark=True)
    section_title(s, "02 / Project", "From monitoring records to decision support", "One reproducible pilot links modelling, evaluation, and a location-focused interface.", True)
    steps = [("01", "Prepare", "Hourly PM₂.₅ + weather\n12 station panel"), ("02", "Forecast", "WAGTF-Net\n+1 · +6 · +24 hours"), ("03", "Communicate", "Hotspots · event probability\nP10–P90 uncertainty")]
    for i, (n, title, body) in enumerate(steps):
        x = .8 + i * 4.2
        add_rect(s, x, 2.4, 3.35, 2.25, "294A43", True)
        add_text(s, n, x+.24, 2.65, .5, .28, 12, AMBER, True)
        add_text(s, title, x+.24, 3.1, 2.9, .42, 22, WHITE, True)
        add_text(s, body, x+.24, 3.7, 2.86, .65, 14, "C4D2CC")
        if i < 2:
            line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+3.45), Inches(3.52), Inches(x+4.0), Inches(3.52))
            line.line.color.rgb = rgb(AMBER); line.line.width = Pt(2)
    add_text(s, "Research focus", .83, 5.34, 1.9, .28, 11, AMBER, True)
    add_text(s, "Benchmark short-horizon forecasts · detect episode risk · test missing-input sensitivity", .83, 5.72, 11.7, .38, 17, WHITE, True)
    add_footer(s, 3, True)

    # 4 · Data
    s = new_slide(prs)
    section_title(s, "03 / Data", "A documented multi-station benchmark", "UCI Beijing Multi-Site Air Quality · historical hourly observations · public CC BY 4.0 dataset")
    add_rect(s, .68, 2.0, 4.0, 4.45, WHITE, True, LINE)
    stats = [("12", "monitoring sites"), ("35,064", "hourly timestamps"), ("420,768", "site-hour records"), ("2013–2017", "benchmark period")]
    for i, (v, label) in enumerate(stats):
        y = 2.25 + i * .79
        add_text(s, v, .98, y, 2.0, .38, 19, GREEN, True)
        add_text(s, label, 3.02, y+.055, 1.28, .26, 10, MUTED)
    add_text(s, "PM₂.₅ · PM₁₀ · SO₂ · NO₂ · CO · O₃\nTemperature · pressure · dew point\nPrecipitation · wind direction · speed", .99, 5.55, 3.4, .62, 9.5, MUTED)
    s.shapes.add_picture(str(RESULTS / "station_map.png"), Inches(4.95), Inches(2.0), width=Inches(5.0), height=Inches(4.45))
    add_rect(s, 10.22, 2.0, 2.4, 4.45, SAGE, True)
    add_text(s, "PILOT SCOPE", 10.48, 2.33, 1.9, .24, 10, GREEN, True)
    add_text(s, "Beijing", 10.48, 2.82, 1.8, .45, 22, INK, True)
    add_text(s, "Public benchmark makes the method reproducible. It does not represent air quality in Kerala or Amritapuri.", 10.48, 3.45, 1.82, 1.65, 13, INK)
    add_text(s, "Site coordinates are approximate values added from published station tables.", 10.48, 5.4, 1.82, .67, 10, MUTED)
    add_footer(s, 4)

    # 5 · Architecture overview
    s = new_slide(prs)
    section_title(s, "04 / Framework", "MOSAIC-Air connects prediction to a readable map", "The pilot is implemented end to end: historical panel → model → API → browser dashboard.")
    blocks = [
        ("INPUT WINDOW", "24 h of pollution\nweather · time\nstation locations", "DCE9DF", GREEN),
        ("WAGTF-NET", "temporal convolution\nwind-biased graph attention\nadaptive fusion", "2D6258", WHITE),
        ("FORECAST", "PM₂.₅ median\nP10 / P90 range\nepisode probability", "EAB45E", INK),
        ("INTERFACE", "station map\nhotspots · weather\nREST API", "D97754", WHITE),
    ]
    for i, (title, body, fill, fg) in enumerate(blocks):
        x=.64+i*3.18
        add_rect(s,x,2.55,2.65,2.05,fill,True)
        add_text(s,title,x+.18,2.83,2.25,.26,10,fg,True,align=PP_ALIGN.CENTER)
        add_text(s,body,x+.18,3.32,2.25,.95,14,fg,False,align=PP_ALIGN.CENTER)
        if i<3:
            l=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x+2.7),Inches(3.56),Inches(x+3.08),Inches(3.56))
            l.line.color.rgb=rgb(GREEN);l.line.width=Pt(2)
    add_rect(s, .78, 5.25, 11.78, .84, PALE, True)
    add_text(s, "Historical replay only", 1.05, 5.52, 2.1, .24, 12, CORAL, True)
    add_text(s, "Current pollution observations and forecast weather are needed before this pipeline can be used live.", 3.08, 5.49, 8.85, .3, 14, INK)
    add_footer(s, 5)

    # 6 · Model
    s = new_slide(prs, dark=True)
    section_title(s, "05 / Model", "Wind changes how station information is shared", "Attention combines learned similarity with direction, speed, and distance.", True)
    # station-to-station conceptual motif
    stations = [(1.2,3.2,"A"),(3.4,2.1,"B"),(3.55,4.5,"C"),(5.7,3.25,"D")]
    for x,y,lbl in stations:
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.56),Inches(.56));c.fill.solid();c.fill.fore_color.rgb=rgb(AMBER);c.line.color.rgb=rgb(WHITE)
        add_text(s,lbl,x+.17,y+.16,.22,.2,11,DARK,True,align=PP_ALIGN.CENTER)
    for a,b in [((1.75,3.45),(3.4,2.4)),((1.75,3.5),(3.55,4.65)),((3.95,2.35),(5.7,3.45)),((4.1,4.65),(5.7,3.55))]:
        l=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));l.line.color.rgb=rgb("66877D");l.line.width=Pt(1.7)
    # wind arrow
    ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(1.0),Inches(5.37),Inches(3.8),Inches(.16));ar.fill.solid();ar.fill.fore_color.rgb=rgb(CORAL);ar.line.fill.background()
    add_text(s,"DOWNWIND",1.0,5.68,1.5,.25,10,"E6B06D",True)
    add_rect(s,6.55,2.1,5.9,3.9,"294A43",True)
    add_text(s,"WAGTF-Net",6.9,2.43,5.15,.42,23,WHITE,True)
    add_bullets(s,["24 h temporal convolution per station","Graph attention across nearby sites","Wind alignment biases source → destination influence","Three horizons × P10 / P50 / P90 + episode probability"],6.9,3.12,5.0,2.35,14,"D7E5DD",12)
    add_text(s,"Approximate wind-informed attention is a learned influence prior, not a physical transport simulator.",6.9,5.42,5.0,.37,10,"B8CCC3")
    add_footer(s, 6, True)

    # 7 · Experimental design
    s = new_slide(prs)
    section_title(s, "06 / Evaluation", "Chronological holdout tests time-forward forecasting", "Inputs use the previous 24 hours; issue times are spaced every six hours.")
    timeline=[("TRAIN","Mar 2013 – Dec 2015",GREEN,WHITE),("VALIDATE","Jan – Jun 2016",AMBER,INK),("TEST","Jul 2016 – Feb 2017",CORAL,WHITE)]
    x=.72
    widths=[5.4,2.75,3.75]
    for i,(name,period,fill,fg) in enumerate(timeline):
        w=widths[i]; add_rect(s,x,2.28,w,1.05,fill,True)
        add_text(s,name,x+.18,2.46,w-.36,.22,10,fg,True)
        add_text(s,period,x+.18,2.78,w-.36,.31,15,fg,True)
        x+=w+.15
    add_text(s,"FORECAST HORIZONS",.78,3.86,2.25,.25,10,GREEN,True)
    for i,h in enumerate(["+1 h","+6 h","+24 h"]):
        x=3.0+i*1.42; add_pill(s,h,x,3.81,1.15,fill=SAGE,color=GREEN)
    add_rect(s,.72,4.55,5.75,1.55,WHITE,True,LINE)
    add_text(s,"BASELINES",.98,4.82,1.3,.23,10,GREEN,True)
    add_text(s,"Persistence · HistGradientBoosting",.98,5.24,5.0,.37,17,INK,True)
    add_rect(s,6.72,4.55,5.88,1.55,WHITE,True,LINE)
    add_text(s,"EVALUATION",6.98,4.82,1.6,.23,10,GREEN,True)
    add_text(s,"MAE · RMSE · R² · episode F1 / Brier · P10–P90 coverage",6.98,5.21,5.22,.5,14,INK,True)
    add_footer(s, 7)

    # 8 · Metrics chart
    s = new_slide(prs)
    section_title(s, "07 / Results", "WAGTF-Net is strongest at six hours in this pilot", "Held-out MAE (µg/m³) · all event cutoffs selected on validation data")
    chart_x, chart_y, chart_w, chart_h = 1.2, 2.3, 8.1, 3.8
    maxv=70
    # horizontal grid lines + labels
    for v in [0,20,40,60]:
        yy=chart_y+chart_h-(v/maxv)*chart_h
        ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(chart_x),Inches(yy),Inches(chart_x+chart_w),Inches(yy));ln.line.color.rgb=rgb(LINE);ln.line.width=Pt(.7)
        add_text(s,str(v),.76,yy-.08,.3,.18,8,MUTED,align=PP_ALIGN.RIGHT)
    colors={"Persistence":MUTED,"HistGradientBoosting":AMBER,"WAGTF-Net":GREEN}
    horizons=[1,6,24]
    base_positions=[2.1,4.75,7.4]
    for hi,h in enumerate(horizons):
        rows=METRICS[(METRICS.horizon_hours==h)&(~METRICS.model.str.contains("missing",case=False))]
        for j,model in enumerate(["Persistence","HistGradientBoosting","WAGTF-Net"]):
            row=rows[rows.model==model].iloc[0];val=float(row.MAE)
            xx=base_positions[hi]+j*.58;bh=(val/maxv)*chart_h;yy=chart_y+chart_h-bh
            add_rect(s,xx,yy,.43,bh,colors[model],False)
            add_text(s,f"{val:.1f}",xx-.06,yy-.27,.55,.2,9,INK,True,align=PP_ALIGN.CENTER)
        add_text(s,f"+{h} h",base_positions[hi]-.02,chart_y+chart_h+.19,1.8,.25,11,INK,True,align=PP_ALIGN.CENTER)
    for j,(name,col) in enumerate(colors.items()):
        add_rect(s,1.4+j*2.0,6.62,.16,.16,col,True)
        add_text(s,name,1.65+j*2.0,6.59,1.75,.22,9,MUTED)
    add_rect(s,9.75,2.27,2.75,3.85,SAGE,True)
    add_text(s,"AT +6 HOURS",10.03,2.62,2.1,.23,10,GREEN,True)
    add_text(s,"28.80",10.03,3.03,2.1,.6,31,INK,True,font="Aptos Display")
    add_text(s,"µg/m³ MAE",10.03,3.65,2.1,.25,11,MUTED)
    add_text(s,"vs 31.19 for Gradient Boosting\n≈ 7.7% lower MAE",10.03,4.27,2.1,.68,13,INK,True)
    add_text(s,"At +1 h, Gradient Boosting has lower MAE (10.07 vs 10.92).",10.03,5.34,2.1,.58,10,MUTED)
    add_footer(s, 8)

    # 9 · Uncertainty/missingness
    s = new_slide(prs, dark=True)
    section_title(s, "08 / Reliability", "Uncertainty helps—but missing inputs still matter", "Nominal 80% P10–P90 coverage is close overall; the long-horizon interval is broad.", True)
    cards=[("+1 h","82.5% coverage","36.7 µg/m³ mean width"),("+6 h","84.1% coverage","95.8 µg/m³ mean width"),("+24 h","80.4% coverage","155.4 µg/m³ mean width")]
    for i,(h,cov,width) in enumerate(cards):
        x=.78+i*4.15;add_rect(s,x,2.22,3.72,1.48,"294A43",True)
        add_text(s,h,x+.23,2.48,.8,.25,11,AMBER,True)
        add_text(s,cov,x+.23,2.85,3.2,.38,22,WHITE,True)
        add_text(s,width,x+.23,3.35,3.2,.24,10,"B8CCC3")
    add_text(s,"ADDITIONAL 20% PM₂.₅ INPUT MISSING · CHANGE IN TEST MAE",.82,4.18,8.4,.24,10,AMBER,True)
    missing=METRICS[METRICS.model.str.contains("20%")].sort_values("horizon_hours")
    standard=METRICS[METRICS.model=="WAGTF-Net"].sort_values("horizon_hours")
    for i,(_,r) in enumerate(missing.iterrows()):
        base=standard.iloc[i]; pct=(float(r.MAE)/float(base.MAE)-1)*100
        x=.92+i*3.95; add_text(s,f"+{int(r.horizon_hours)} h",x,4.67,1,.2,10,"B8CCC3",True)
        add_text(s,f"+{pct:.1f}%",x,5.0,2.3,.48,25,CORAL if pct>30 else AMBER,True)
        add_text(s,f"{float(base.MAE):.1f} → {float(r.MAE):.1f} MAE",x,5.57,2.7,.25,11,WHITE)
    add_text(s,"Randomly masking extra input cells without retraining. One-hour forecasts degrade most; add missing-data augmentation before deployment.",.82,6.33,11.4,.35,10,"B8CCC3")
    add_footer(s, 9, True)

    # 10 · Prototype
    s = new_slide(prs)
    section_title(s, "09 / Prototype", "The dashboard makes forecasts location-specific", "A local REST API serves the selected historical issue time to the browser interface.")
    # map image + mock UI cards
    add_rect(s,.68,2.1,7.15,4.15,WHITE,True,LINE)
    add_text(s,"STATION MAP · GEOGRAPHIC CONTEXT",.95,2.39,5.9,.24,10,GREEN,True)
    s.shapes.add_picture(str(RESULTS / "station_map.png"),Inches(.95),Inches(2.78),width=Inches(4.45),height=Inches(3.12))
    add_rect(s,5.66,2.75,1.83,3.2,PALE,True)
    add_text(s,"SELECTED SITE",5.85,3.02,1.45,.2,9,GREEN,True)
    add_text(s,"Station detail",5.85,3.4,1.45,.34,15,INK,True,font="Aptos Display")
    add_text(s,"PM₂.₅ median",5.85,4.05,1.4,.2,9,MUTED)
    add_text(s,"P50",5.85,4.3,1.3,.42,23,INK,True)
    add_text(s,"P10–P90 range\nEpisode probability\nWind · temperature",5.85,4.95,1.47,.66,9,MUTED)
    add_text(s,"INCLUDED",8.35,2.35,1.5,.24,10,CORAL,True)
    add_bullets(s,["Interactive station map","1 / 6 / 24 h horizon selector","P10–P90 uncertainty range","Episode probability + hotspot ranking","FastAPI JSON endpoints"],8.35,2.85,4.0,2.7,14,INK,13)
    add_rect(s,8.34,5.55,4.15,.68,SAGE,True)
    add_text(s,"Historical replay · not live · not a health advisory",8.55,5.75,3.75,.22,10,GREEN,True)
    add_footer(s, 10)

    # 11 · Caveats
    s = new_slide(prs)
    section_title(s, "10 / Interpretation", "What this pilot can—and cannot—say", "The result is promising at the medium horizon and bounded by its benchmark design.")
    add_rect(s,.72,2.18,5.77,3.82,SAGE,True)
    add_text(s,"WHAT WE FOUND",1.02,2.5,2.2,.24,10,GREEN,True)
    add_bullets(s,["6 h MAE improves about 7.7% over HistGradientBoosting","24 h improvement is modest; intervals are wide","One-hour forecasts depend on recent PM₂.₅ inputs","Forecasts and hotspot views work end to end"],1.02,2.97,5.0,2.55,14,INK,12)
    add_rect(s,6.78,2.18,5.78,3.82,WHITE,True,LINE)
    add_text(s,"NEXT VALIDATION",7.08,2.5,2.6,.24,10,CORAL,True)
    add_bullets(s,["Use a local dataset; test held-out stations and cities","Join live observations to forecast weather","Stress-test outages, drift, and systematic missingness","Calibrate by station/horizon; review thresholds locally"],7.08,2.97,5.0,2.55,14,INK,12)
    add_text(s,"The 75 µg/m³ label is a project benchmark—not a health or regulatory threshold.",.82,6.32,11.4,.28,11,MUTED,True)
    add_footer(s, 11)

    # 12 · Close
    s = new_slide(prs, dark=True)
    add_pill(s,"MOSAIC-AIR · PILOT CONCLUSION",.78,.7,3.3,fill="31564D",color="D5E6DB")
    add_text(s,"A reproducible path from\nstation data to spatial forecasts.",.8,1.58,8.2,1.55,33,WHITE,True,font="Aptos Display")
    add_text(s,"WAGTF-Net shows value at 6 hours in this Beijing case study. The next step is validation with local data, weather forecasts, and environmental stakeholders.",.83,3.65,7.8,.9,16,"C4D2CC")
    add_rect(s,.84,5.23,3.25,.72,"294A43",True)
    add_text(s,"Thank you",1.08,5.46,2.75,.28,17,AMBER,True)
    # compact network motif
    for x,y in [(10.0,2.0),(11.5,1.45),(12.0,3.1),(10.8,4.2),(9.5,4.0),(11.4,5.55)]:
        c=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.19),Inches(.19));c.fill.solid();c.fill.fore_color.rgb=rgb(AMBER);c.line.fill.background()
    for a,b in [((10.1,2.1),(11.6,1.55)),((11.6,1.55),(12.1,3.2)),((12.1,3.2),(10.9,4.3)),((10.9,4.3),(9.6,4.1)),((10.9,4.3),(11.5,5.65)),((10.1,2.1),(10.9,4.3))]:
        l=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));l.line.color.rgb=rgb("52746A");l.line.width=Pt(1.4)
    add_footer(s, 12, True)
    out=ROOT / "MOSAIC-Air_Presentation.pptx"
    prs.save(out)
    return out


def _inline_runs(paragraph, text):
    # Preserve simple bold and inline-code markup in the Markdown report.
    parts=re.split(r"(\*\*.*?\*\*|`[^`]+`)",text)
    for part in parts:
        if not part: continue
        run=paragraph.add_run(part[2:-2] if part.startswith("**") else part[1:-1] if part.startswith("`") else part)
        run.bold=part.startswith("**")
        if part.startswith("`"):
            run.font.name="Consolas";run.font.size=DocPt(9)


def markdown_to_docx(md_path, out_path):
    doc=Document()
    section=doc.sections[0]
    section.top_margin=DocInches(.72);section.bottom_margin=DocInches(.72);section.left_margin=DocInches(.82);section.right_margin=DocInches(.82)
    styles=doc.styles
    styles["Normal"].font.name="Aptos";styles["Normal"].font.size=DocPt(10)
    styles["Normal"].font.color.rgb=DocRGBColor.from_string(INK)
    for level,size in [(1,21),(2,15),(3,12)]:
        st=styles[f"Heading {level}"];st.font.name="Aptos Display";st.font.size=DocPt(size);st.font.bold=True;st.font.color.rgb=DocRGBColor.from_string(GREEN)
    lines=md_path.read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line.startswith("!["):
            m=re.search(r"\((.*?)\)",line)
            if m:
                img=(md_path.parent/m.group(1)).resolve()
                if img.exists():
                    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(img),width=DocInches(6.25))
            i+=1;continue
        if line.startswith("|"):
            table_lines=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip());i+=1
            rows=[]
            for tl in table_lines:
                cells=[c.strip() for c in tl.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?",c) for c in cells): continue
                rows.append(cells)
            if rows:
                cols=max(map(len,rows));tab=doc.add_table(rows=0,cols=cols);tab.style="Light Shading Accent 1"
                for ri,row in enumerate(rows):
                    cells=tab.add_row().cells
                    for ci in range(cols):
                        val=row[ci] if ci<len(row) else ""
                        cells[ci].text=val.replace("**","")
                        for p in cells[ci].paragraphs:
                            for r in p.runs:
                                r.font.name="Aptos";r.font.size=DocPt(8.5)
                                if ri==0:r.bold=True
                doc.add_paragraph()
            continue
        if line.startswith("#"):
            level=min(len(line)-len(line.lstrip("#")),3)
            title=line[level:].strip()
            p=doc.add_heading(level=level);_inline_runs(p,title)
            i+=1;continue
        if line.startswith("- ") or re.match(r"\d+\. ",line):
            ordered=bool(re.match(r"\d+\. ",line));content=re.sub(r"^(?:- |\d+\. )","",line)
            if ordered:
                p=doc.add_paragraph();_inline_runs(p,line)
            else:
                p=doc.add_paragraph(style="List Bullet");_inline_runs(p,content)
            i+=1;continue
        p=doc.add_paragraph();_inline_runs(p,line)
        p.paragraph_format.space_after=DocPt(5)
        i+=1
    # Add a consistent footer.
    for sec in doc.sections:
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("MOSAIC-Air  ·  Environmental Management");r.font.size=DocPt(8);r.font.color.rgb=DocRGBColor.from_string(MUTED)
    doc.save(out_path)


def export_pdf(path):
    subprocess.run(["libreoffice","--headless","--convert-to","pdf","--outdir",str(path.parent),str(path)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)


def main():
    pptx=make_slides()
    report=DOCS/"project_report.docx"
    markdown_to_docx(DOCS/"project_report.md",report)
    export_pdf(report)
    export_pdf(pptx)
    print(f"Created {pptx.relative_to(ROOT)} and {report.relative_to(ROOT)} (+ PDFs)")


if __name__=="__main__":
    main()
